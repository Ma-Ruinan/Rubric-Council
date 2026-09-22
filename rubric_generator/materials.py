from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from .dataset import PROBLEM_FILENAME, is_reference_rubric


TEXT_SUFFIXES = {".txt", ".md", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml", ".xml", ".html", ".htm"}


def prepare_material_evidence(task_dir: Path, output_root: Path) -> Path:
    """Create a private, machine-readable source inventory for solve-to-audit work."""
    output_root.mkdir(parents=True, exist_ok=True)
    extracted = output_root / "extracted"
    extracted.mkdir(parents=True, exist_ok=True)
    entries: list[dict[str, Any]] = []
    for source in sorted(task_dir.rglob("*"), key=lambda path: path.as_posix().casefold()):
        if not source.is_file() or source.name == PROBLEM_FILENAME or is_reference_rubric(source):
            continue
        relative = source.relative_to(task_dir)
        entry: dict[str, Any] = {
            "path": relative.as_posix(),
            "size_bytes": source.stat().st_size,
            "sha256": _sha256(source),
            "suffix": source.suffix.casefold(),
        }
        try:
            details, artifacts = _extract(source, relative, extracted)
            entry["inspection"] = details
            entry["extracted_artifacts"] = [path.relative_to(output_root).as_posix() for path in artifacts]
            entry["status"] = "extracted"
        except Exception as exc:  # Preserve the file and make the limitation explicit.
            entry["status"] = "inspection_failed"
            entry["error"] = f"{type(exc).__name__}: {exc}"
        entries.append(entry)
    manifest = {
        "task_dir": task_dir.as_posix(),
        "problem_file": (task_dir / PROBLEM_FILENAME).as_posix(),
        "materials": entries,
        "notice": "Extraction is evidence access, not gold verification. Anchors still require recorded recomputation and cross-checks.",
    }
    path = output_root / "material-manifest.json"
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return path


def _extract(source: Path, relative: Path, output: Path) -> tuple[dict[str, Any], list[Path]]:
    suffix = source.suffix.casefold()
    if suffix in TEXT_SUFFIXES:
        return _extract_text(source, relative, output)
    if suffix == ".docx":
        return _extract_docx(source, relative, output)
    if suffix == ".pptx":
        return _extract_pptx(source, relative, output)
    if suffix == ".xlsx":
        return _extract_xlsx(source, relative, output)
    if suffix == ".pdf":
        return _extract_pdf(source, relative, output)
    return {"kind": "binary", "note": "No structured extractor configured"}, []


def _output_path(relative: Path, output: Path, suffix: str) -> Path:
    slug = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff._-]+", "-", relative.as_posix()).strip("-.")
    digest = hashlib.sha256(relative.as_posix().encode("utf-8")).hexdigest()[:8]
    return output / f"{slug[:80]}-{digest}{suffix}"


def _extract_text(source: Path, relative: Path, output: Path) -> tuple[dict[str, Any], list[Path]]:
    text = source.read_text(encoding="utf-8-sig", errors="replace")
    target = _output_path(relative, output, ".txt")
    target.write_text(text, encoding="utf-8")
    return {"kind": "text", "characters": len(text), "lines": len(text.splitlines())}, [target]


def _extract_docx(source: Path, relative: Path, output: Path) -> tuple[dict[str, Any], list[Path]]:
    from docx import Document

    document = Document(source)
    lines = [f"P{i}: {paragraph.text}" for i, paragraph in enumerate(document.paragraphs, 1) if paragraph.text.strip()]
    for table_index, table in enumerate(document.tables, 1):
        lines.append(f"TABLE {table_index}")
        for row_index, row in enumerate(table.rows, 1):
            lines.append(f"T{table_index}R{row_index}: " + " | ".join(cell.text for cell in row.cells))
    target = _output_path(relative, output, ".txt")
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return {"kind": "docx", "paragraphs": len(document.paragraphs), "tables": len(document.tables)}, [target]


def _extract_pptx(source: Path, relative: Path, output: Path) -> tuple[dict[str, Any], list[Path]]:
    from pptx import Presentation

    presentation = Presentation(source)
    lines: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, 1):
        lines.append(f"SLIDE {slide_index}")
        for shape_index, shape in enumerate(slide.shapes, 1):
            if hasattr(shape, "text") and shape.text.strip():
                lines.append(f"S{slide_index}O{shape_index}: {shape.text}")
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows, 1):
                    lines.append(f"S{slide_index}O{shape_index}R{row_index}: " + " | ".join(cell.text for cell in row.cells))
    target = _output_path(relative, output, ".txt")
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return {"kind": "pptx", "slides": len(presentation.slides)}, [target]


def _extract_pdf(source: Path, relative: Path, output: Path) -> tuple[dict[str, Any], list[Path]]:
    from pypdf import PdfReader

    reader = PdfReader(source)
    lines: list[str] = []
    extracted_pages = 0
    for page_index, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            extracted_pages += 1
        lines.append(f"===== PDF PAGE {page_index} =====\n{text}")
    target = _output_path(relative, output, ".txt")
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return {"kind": "pdf", "pages": len(reader.pages), "pages_with_text": extracted_pages}, [target]


def _extract_xlsx(source: Path, relative: Path, output: Path) -> tuple[dict[str, Any], list[Path]]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    formulas = load_workbook(source, read_only=True, data_only=False)
    cached = load_workbook(source, read_only=True, data_only=True)
    target = _output_path(relative, output, ".jsonl")
    sheets: list[dict[str, Any]] = []
    with target.open("w", encoding="utf-8", newline="\n") as handle:
        for sheet_name in formulas.sheetnames:
            sheet = formulas[sheet_name]
            cached_sheet = cached[sheet_name]
            nonempty = 0
            formula_count = 0
            duplicate_rows = 0
            seen_rows: Counter[tuple[str, ...]] = Counter()
            column_nonempty: Counter[str] = Counter()
            for row_index, (formula_row, cached_row) in enumerate(zip(sheet.iter_rows(), cached_sheet.iter_rows()), 1):
                row_values: list[str] = []
                cells: list[dict[str, Any]] = []
                for column_index, (formula_cell, cached_cell) in enumerate(zip(formula_row, cached_row), 1):
                    value = formula_cell.value
                    cached_value = cached_cell.value
                    column_letter = get_column_letter(column_index)
                    coordinate = f"{column_letter}{row_index}"
                    if value is not None:
                        nonempty += 1
                        column_nonempty[column_letter] += 1
                    if formula_cell.data_type == "f":
                        formula_count += 1
                    row_values.append(_stable_value(value))
                    if value is not None or cached_value is not None:
                        cells.append({
                            "cell": coordinate,
                            "value": value,
                            "cached_value": cached_value,
                            "data_type": formula_cell.data_type,
                            "number_format": formula_cell.number_format,
                        })
                key = tuple(row_values)
                if any(part for part in key):
                    if seen_rows[key]:
                        duplicate_rows += 1
                    seen_rows[key] += 1
                if cells:
                    handle.write(json.dumps({"sheet": sheet_name, "row": row_index, "cells": cells}, ensure_ascii=False, default=str) + "\n")
            sheets.append({
                "name": sheet_name,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "nonempty_cells": nonempty,
                "formula_cells": formula_count,
                "extra_exact_duplicate_rows": duplicate_rows,
                "column_nonempty": dict(sorted(column_nonempty.items())),
            })
    formulas.close()
    cached.close()
    return {"kind": "xlsx", "sheets": sheets}, [target]


def _stable_value(value: Any) -> str:
    if value is None:
        return ""
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
